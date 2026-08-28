from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import unittest

from docx import Document

from atlas_datagob.services.business_case_document import (
    DOCX_CONTENT_TYPE,
    build_business_case_docx,
    business_case_document_filename,
)


class BusinessCaseDocumentTest(unittest.TestCase):
    def setUp(self) -> None:
        self.business_case = {
            "business_problem": "Comercial detecta quiebres de stock demasiado tarde.",
            "desired_outcome": "Anticipar el riesgo de quiebre de stock.",
            "business_area": "Planificación Comercial",
            "impacted_process": "Planificación y reposición de inventario",
            "current_situation": "La revisión actual es reactiva.",
            "stakeholders": ["Planificación Comercial", "Comercial"],
            "success_metrics": ["Reducir quiebres de stock al menos 20%"],
            "data_sources": ["SAP ventas", "SAP inventarios"],
            "project_classification": {
                "primary_type": "machine_learning",
                "subtype": "forecasting",
                "secondary_capabilities": ["data_engineering", "dashboard_analytics"],
                "confidence": 0.92,
            },
            "data_readiness": {"score": 100, "status": "ready", "gaps": []},
            "architecture_assessment": {
                "pattern_id": "GCP-ML-001",
                "pattern_version": "1.0",
                "pattern_name": "Governed ML lifecycle",
                "status": "approved_baseline",
                "gcp_services": ["Cloud Storage", "BigQuery", "Vertex AI"],
                "required_components": ["Ingestion", "Model serving", "Monitoring"],
                "human_architecture_review_required": False,
            },
            "policy_assessment": {
                "status": "requirements_identified",
                "policy_references": ["DATA-001@1.0", "ML-001@1.0"],
                "missing_controls": ["model_registry"],
            },
            "preliminary_risk": "medium",
            "definition_gaps": [],
            "governance_requirements": ["model_registry", "drift_monitoring"],
            "recommendation": "Continuar a Comité Operativo.",
            "completeness": 100,
            "ready_to_register": True,
        }

    def test_docx_is_valid_and_contains_canonical_sections(self) -> None:
        payload = build_business_case_docx(
            self.business_case,
            session_id="intake-test123456",
            generated_at=datetime(2026, 8, 27, 12, 0, tzinfo=timezone.utc),
        )

        self.assertGreater(len(payload), 10_000)
        document = Document(BytesIO(payload))
        paragraphs = "\n".join(paragraph.text for paragraph in document.paragraphs)
        table_text = "\n".join(
            cell.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
        )
        content = f"{paragraphs}\n{table_text}"

        self.assertIn("Caso de Negocio", content)
        self.assertIn("Resumen ejecutivo", content)
        self.assertIn("Comercial detecta quiebres de stock demasiado tarde.", content)
        self.assertIn("GCP-ML-001", content)
        self.assertIn("DATA-001@1.0", content)
        self.assertIn("Definition of Ready", content)
        self.assertIn("Cumplida", content)
        self.assertEqual(document.core_properties.author, "ATLAS DataGob")

    def test_filename_is_safe_and_docx_content_type_is_official(self) -> None:
        self.assertEqual(
            business_case_document_filename("intake-test/unsafe:123"),
            "ATLAS_Caso_de_Negocio_intake-testunsafe123.docx",
        )
        self.assertEqual(
            DOCX_CONTENT_TYPE,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )


if __name__ == "__main__":
    unittest.main()
