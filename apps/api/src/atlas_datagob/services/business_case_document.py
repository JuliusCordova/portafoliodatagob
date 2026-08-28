from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from typing import Any, Iterable

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

DOCX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

PURPLE = "6D46E8"
DARK = "17171F"
MUTED = "666675"
LIGHT = "F7F5FC"
LINE = "E6E3EE"
GREEN = "1F9D64"
AMBER = "B7791F"


def _text(value: Any, fallback: str = "No definido") -> str:
    if value is None:
        return fallback
    if isinstance(value, str):
        value = value.strip()
        return value or fallback
    return str(value)


def _items(value: Any) -> list[str]:
    if not isinstance(value, list):
        return []
    return [str(item).strip() for item in value if str(item).strip()]


def _nested(data: dict[str, Any], *keys: str) -> dict[str, Any]:
    current: Any = data
    for key in keys:
        if not isinstance(current, dict):
            return {}
        current = current.get(key)
    return current if isinstance(current, dict) else {}


def _set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def _set_cell_margins(cell, top=90, start=120, bottom=90, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_cell_border(cell, color: str = LINE, size: int = 6) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), color)


def _style_table(table) -> None:
    table.autofit = True
    for row in table.rows:
        row_properties = row._tr.get_or_add_trPr()
        if row_properties.find(qn("w:cantSplit")) is None:
            row_properties.append(OxmlElement("w:cantSplit"))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            _set_cell_margins(cell)
            _set_cell_border(cell)
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(9)


def _add_key_value_table(doc: Document, rows: Iterable[tuple[str, str]]) -> None:
    values = list(rows)
    table = doc.add_table(rows=len(values), cols=2)
    table.columns[0].width = Inches(1.85)
    table.columns[1].width = Inches(4.95)
    for index, (label, value) in enumerate(values):
        left, right = table.rows[index].cells
        _set_cell_shading(left, LIGHT)
        left.paragraphs[0].add_run(label).bold = True
        right.paragraphs[0].add_run(value)
    _style_table(table)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def _add_bullets(doc: Document, values: list[str], empty: str = "No definido") -> None:
    if not values:
        doc.add_paragraph(empty)
        return
    for value in values:
        paragraph = doc.add_paragraph(value, style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)


def _add_section_heading(doc: Document, number: str, title: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(f"{number}. {title}")
    run.bold = True
    run.font.size = Pt(13.5)
    run.font.color.rgb = RGBColor.from_string(DARK)


def _add_status_badge(cell, text: str, fill: str) -> None:
    _set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)


def build_business_case_docx(
    business_case: dict[str, Any],
    *,
    session_id: str,
    generated_at: datetime | None = None,
) -> bytes:
    """Build a deterministic executive DOCX from the canonical Business Case only."""

    generated_at = generated_at or datetime.now(timezone.utc)
    classification = _nested(business_case, "project_classification")
    readiness = _nested(business_case, "data_readiness")
    architecture = _nested(business_case, "architecture_assessment")
    policy = _nested(business_case, "policy_assessment")

    doc = Document()
    doc.core_properties.title = "ATLAS DataGob · Caso de Negocio"
    doc.core_properties.subject = "Canonical Business Case generado desde Intake Conversacional Gobernado"
    doc.core_properties.author = "ATLAS DataGob"
    doc.core_properties.keywords = "business case, data governance, AI governance, ATLAS DataGob"

    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(DARK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name in ("Title", "Heading 1", "Heading 2", "Heading 3"):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor.from_string(DARK)

    header_paragraph = section.header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_run = header_paragraph.add_run("ATLAS DataGob · Caso de Negocio")
    header_run.bold = True
    header_run.font.name = "Arial"
    header_run.font.size = Pt(8)
    header_run.font.color.rgb = RGBColor.from_string(PURPLE)

    footer_paragraph = section.footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_paragraph.add_run(
        "Generado desde Intake Conversacional Gobernado · Documento de trabajo sujeto a revisión"
    )
    footer_run.font.name = "Arial"
    footer_run.font.size = Pt(7.5)
    footer_run.font.color.rgb = RGBColor.from_string(MUTED)

    brand = doc.add_paragraph()
    brand.paragraph_format.space_before = Pt(34)
    brand.paragraph_format.space_after = Pt(7)
    brand_run = brand.add_run("ATLAS DataGob")
    brand_run.bold = True
    brand_run.font.name = "Arial"
    brand_run.font.size = Pt(12)
    brand_run.font.color.rgb = RGBColor.from_string(PURPLE)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("Caso de Negocio")
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(27)
    title_run.font.color.rgb = RGBColor.from_string(DARK)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    subtitle_run = subtitle.add_run(
        f"{_text(business_case.get('business_area'))} · "
        f"{_text(classification.get('primary_type')).replace('_', ' ').title()}"
    )
    subtitle_run.font.name = "Arial"
    subtitle_run.font.size = Pt(12)
    subtitle_run.font.color.rgb = RGBColor.from_string(MUTED)

    completeness = int(business_case.get("completeness") or 0)
    ready = bool(business_case.get("ready_to_register"))
    status_table = doc.add_table(rows=1, cols=3)
    cells = status_table.rows[0].cells
    cells[0].paragraphs[0].add_run(f"{completeness}%\nDefinición").bold = True
    _add_status_badge(
        cells[1],
        "LISTO PARA REGISTRO" if ready else "EN DEFINICIÓN",
        GREEN if ready else AMBER,
    )
    cells[2].paragraphs[0].add_run(f"Sesión\n{session_id}").bold = True
    _style_table(status_table)

    doc.add_paragraph()
    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(16)
    note_run = note.add_run(
        "Este documento se genera de forma determinística desde el Canonical Business Case de ATLAS. "
        "No constituye aprobación del Comité, autorización de inversión ni pase a producción."
    )
    note_run.italic = True
    note_run.font.color.rgb = RGBColor.from_string(MUTED)

    _add_key_value_table(
        doc,
        [
            ("Generado por", "ATLAS DataGob"),
            ("Fecha de generación", generated_at.astimezone(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
            ("Tipo de iniciativa", _text(classification.get("primary_type")).replace("_", " ").title()),
            ("Subtipo", _text(classification.get("subtype")).replace("_", " ").title()),
            ("Riesgo preliminar", _text(business_case.get("preliminary_risk")).replace("_", " ").title()),
        ],
    )

    _add_section_heading(doc, "1", "Resumen ejecutivo")
    summary = doc.add_paragraph()
    summary.add_run("Problema. ").bold = True
    summary.add_run(_text(business_case.get("business_problem")))
    summary.add_run("\nResultado esperado. ").bold = True
    summary.add_run(_text(business_case.get("desired_outcome")))
    summary.add_run("\nRecomendación ATLAS. ").bold = True
    summary.add_run(
        _text(
            business_case.get("recommendation"),
            "Continuar con el flujo gobernado de evaluación.",
        )
    )

    _add_section_heading(doc, "2", "Problema u oportunidad de negocio")
    doc.add_paragraph(_text(business_case.get("business_problem")))
    _add_key_value_table(
        doc,
        [
            ("Situación actual", _text(business_case.get("current_situation"))),
            ("Proceso impactado", _text(business_case.get("impacted_process"))),
        ],
    )

    _add_section_heading(doc, "3", "Objetivo y resultado esperado")
    doc.add_paragraph(_text(business_case.get("desired_outcome")))
    doc.add_paragraph("Criterios de éxito / KPIs:").runs[0].bold = True
    _add_bullets(doc, _items(business_case.get("success_metrics")))

    doc.add_page_break()
    _add_section_heading(doc, "4", "Stakeholders y alcance organizacional")
    _add_key_value_table(
        doc,
        [
            ("Área responsable", _text(business_case.get("business_area"))),
            ("Proceso", _text(business_case.get("impacted_process"))),
        ],
    )
    doc.add_paragraph("Stakeholders identificados:").runs[0].bold = True
    _add_bullets(doc, _items(business_case.get("stakeholders")))

    _add_section_heading(doc, "5", "Datos y Data Readiness")
    _add_key_value_table(
        doc,
        [
            (
                "Readiness score",
                f"{readiness.get('score')}%" if readiness.get("score") is not None else "No evaluado",
            ),
            ("Estado", _text(readiness.get("status")).replace("_", " ").title()),
        ],
    )
    doc.add_paragraph("Fuentes de datos:").runs[0].bold = True
    _add_bullets(doc, _items(business_case.get("data_sources")))
    doc.add_paragraph("Brechas de datos:").runs[0].bold = True
    _add_bullets(
        doc,
        _items(readiness.get("gaps")),
        "Sin brechas de Data Readiness registradas.",
    )

    _add_section_heading(doc, "6", "Clasificación de la iniciativa")
    _add_key_value_table(
        doc,
        [
            ("Tipo principal", _text(classification.get("primary_type")).replace("_", " ").title()),
            ("Subtipo", _text(classification.get("subtype")).replace("_", " ").title()),
            ("Tipo de agente", _text(classification.get("agent_type")).replace("_", " ").title()),
            (
                "Confianza",
                f"{float(classification.get('confidence')) * 100:.0f}%"
                if isinstance(classification.get("confidence"), (int, float))
                else "No disponible",
            ),
        ],
    )
    doc.add_paragraph("Capacidades secundarias:").runs[0].bold = True
    _add_bullets(
        doc,
        [item.replace("_", " ").title() for item in _items(classification.get("secondary_capabilities"))],
    )

    _add_section_heading(doc, "7", "Arquitectura GCP recomendada")
    _add_key_value_table(
        doc,
        [
            ("Patrón", _text(architecture.get("pattern_id"))),
            ("Versión", _text(architecture.get("pattern_version"))),
            ("Nombre", _text(architecture.get("pattern_name"))),
            ("Estado", _text(architecture.get("status")).replace("_", " ").title()),
            (
                "Revisión humana",
                "Sí" if architecture.get("human_architecture_review_required") else "No",
            ),
        ],
    )
    doc.add_paragraph("Servicios GCP identificados:").runs[0].bold = True
    _add_bullets(doc, _items(architecture.get("gcp_services")))
    doc.add_paragraph("Componentes requeridos:").runs[0].bold = True
    _add_bullets(doc, _items(architecture.get("required_components")))

    _add_section_heading(doc, "8", "Gobierno, políticas y controles")
    policy_references = _items(policy.get("policy_references"))
    _add_key_value_table(
        doc,
        [
            ("Estado de políticas", _text(policy.get("status")).replace("_", " ").title()),
            (
                "Políticas aplicables",
                ", ".join(policy_references) if policy_references else "No evaluado",
            ),
        ],
    )
    doc.add_paragraph("Controles/requisitos de gobierno pendientes:").runs[0].bold = True
    governance_requirements = _items(business_case.get("governance_requirements"))
    if not governance_requirements:
        governance_requirements = _items(policy.get("missing_controls"))
    _add_bullets(doc, governance_requirements, "Sin requisitos pendientes registrados.")

    _add_section_heading(doc, "9", "Riesgos, brechas y consideraciones")
    _add_key_value_table(
        doc,
        [
            (
                "Riesgo preliminar",
                _text(business_case.get("preliminary_risk")).replace("_", " ").title(),
            )
        ],
    )
    doc.add_paragraph("Brechas de definición:").runs[0].bold = True
    definition_gaps = _items(business_case.get("definition_gaps")) or _items(business_case.get("gaps"))
    _add_bullets(doc, definition_gaps, "Sin brechas de definición abiertas.")

    _add_section_heading(doc, "10", "Próximos pasos")
    _add_bullets(
        doc,
        [
            "Confirmar que el Caso de Negocio representa correctamente la necesidad.",
            "Registrar el requerimiento formal en ATLAS DataGob.",
            "Continuar con Comité Operativo, scoring y priorización del portafolio.",
            "Resolver los requisitos de gobierno y controles en diseño/delivery antes de producción.",
        ],
    )

    _add_section_heading(doc, "11", "Trazabilidad del artefacto")
    _add_key_value_table(
        doc,
        [
            ("Session ID", session_id),
            ("Completeness", f"{completeness}%"),
            ("Definition of Ready", "Cumplida" if ready else "Pendiente"),
            ("Fuente", "Canonical Business Case · ATLAS Conversational Intake"),
        ],
    )

    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def business_case_document_filename(session_id: str) -> str:
    safe = "".join(ch for ch in session_id if ch.isalnum() or ch in {"-", "_"})[:63]
    return f"ATLAS_Caso_de_Negocio_{safe or 'session'}.docx"
